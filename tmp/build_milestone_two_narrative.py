from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

output = Path('output/doc/Payton_Castle_CS499_Milestone_Two_Narrative.docx')
output.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
for name in ('Normal', 'Title', 'Heading 1'):
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        style.element.rPr.rFonts.attrib.pop(qn('w:' + attr), None)
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
doc.styles['Normal'].paragraph_format.first_line_indent = Inches(.5)
doc.styles['Normal'].paragraph_format.widow_control = True
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.first_line_indent = Inches(0)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
title = doc.add_paragraph('Clearspace Software Engineering Enhancement Narrative', 'Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for line in ('Payton Castle', 'CS 499 Milestone Two', 'September 15, 2026'):
    p = doc.add_paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)

sections = [
('Artifact Background', [
"For this milestone, I enhanced Clearspace by separating responsibilities, improving file-operation feedback, and adding automated tests. I created Clearspace in August 2026 as a Windows file-management application built with C# and Windows Presentation Foundation. It supports folder navigation, file searches, tags, and operations such as copying, moving, renaming, and deleting files. This enhancement focuses on software engineering and design while keeping the application's existing features."
]),
('Reason for Selection and Improvements', [
"I selected Clearspace for my ePortfolio because it provides a practical example of maintaining software with several connected features. Its interface, background searches, and Windows file operations create opportunities to demonstrate design decisions beyond adding a new button or screen. The original project already separated models, commands, and services, but MainViewModel coordinated too many responsibilities. I moved search coordination into SearchCoordinator, navigation loading into NavigationCoordinator, and sidebar state into SidebarViewModel. These components make individual behaviors easier to inspect and test without constructing the entire interface.",
"I also replaced the true-or-false results from copy, move, rename, and delete with FileOperationResult. The new result distinguishes success, cancellation, and failure, while retaining Windows diagnostic information and a message for the user. For example, canceling a copy operation should not be presented as a permission error. Updating the callers was an important part of this change because collecting better information would have little value if the interface ignored it.",
"The enhanced artifact includes 87 automated tests that passed in the Release configuration. These tests check search cancellation and duplicate removal, tag persistence and assignment behavior, indexing edge cases, and file-operation outcomes. Controlled test data makes it possible to simulate failures without deleting or moving personal files. The tests provide repeatable evidence that these specific behaviors work, although they do not prove that every interface interaction is correct."
]),
('Course Outcomes and Enhancement Scope', [
"The completed changes meet the software engineering goals described in my code review: smaller components, more informative operation results, and automated tests. They provide evidence for the outcome concerning well-founded computing techniques and tools by applying separation of responsibilities and repeatable testing to an existing application. The review, implementation notes, and this narrative also support professional technical communication by explaining the reasons for the changes and their limitations.",
"Input validation and explicit handling of cancellation and errors support progress toward a security mindset. These improvements help avoid misleading feedback when files are involved, but they do not guarantee that a partially completed Windows operation can be undone. My enhancement scope remains focused on software engineering. Changes to traversal depth, folder-size calculations, and SQLite tag storage remain planned for the later algorithms and database categories."
]),
('Reflection on Learning and Challenges', [
"This work reinforced the importance of defining what each component owns before moving code. Search was especially challenging because several sources can return results at different times. Canceling an old search is not sufficient unless late results are also prevented from replacing a newer search. Separating the coordinator from its sources made these situations easier to reproduce with controlled tests.",
"Testing also exposed assumptions that were easy to miss during a code review. A filter-only search such as ext:txt could skip subfolder results because the application treated index coverage as proof that the index could answer the query. Clearing a search also restored the folder's items without restoring its status message. Both issues were corrected and covered by regression tests. This experience helped me understand why successful compilation is only one part of verification. Manual checks remain necessary for Windows dialogs and visual behavior, and the latest search fixes have not yet been rechecked in the interface."
])]
for heading, paragraphs in sections:
    doc.add_paragraph(heading, 'Heading 1')
    for text in paragraphs:
        doc.add_paragraph(text)
doc.core_properties.author = 'Payton Castle'
doc.core_properties.title = 'Clearspace Software Engineering Enhancement Narrative'
doc.core_properties.subject = 'CS 499 Milestone Two'
doc.save(output)
print(output.resolve())
print('Body words:', sum(len(p.split()) for _, ps in sections for p in ps))
