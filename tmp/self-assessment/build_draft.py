from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path('output/doc/Payton_Castle_CS499_Professional_Self_Assessment_Draft.docx')
out.parent.mkdir(parents=True, exist_ok=True)
sections = [
('Professional Growth and Career Direction', [
"I entered the Computer Science program because I wanted to make software and games. My coursework and the CS 499 capstone helped me understand the work required to make an application dependable and useful. I now consider how software behaves when an operation fails, how it handles large collections of data, and how clearly it explains its behavior to users. This portfolio introduces my development in those areas through Clearspace, a Windows file manager I created in C# and Windows Presentation Foundation.",
"Developing the ePortfolio helped me evaluate my work through evidence. Instead of describing a feature only by what it does, I explain the problem it addresses, the design choices behind it, and the checks used to verify it. The three milestone snapshots document 87, 158, and 251 passing Release tests, respectively. These are cumulative suite totals at different stages, rather than separate counts to add together. They support specific behaviors in the submitted versions and do not establish that every later change or every Windows environment has been verified.",
"My professional interests remain software and game development. I want to build applications that people can understand and trust, while continuing to explore Unity and real-time interaction. Building VRChat worlds as a hobby gives me another setting in which to develop these interests. My capstone strengthened the habits I want to bring to professional work: investigating problems, testing changes, explaining limitations, and treating users' files as valuable data."
]),
('Collaboration and Communication', [
"I developed Clearspace independently and did not work with other students on a course project. My experience collaborating on Unity projects came before this course. I served as the lead developer and handled the development work, while another contributor provided music. This experience gives me a starting point for working across different creative and technical responsibilities. It also helps me recognize that a software project can include contributions beyond programming.",
"For Outcome 1, building collaborative environments that support decision making, my Unity experience is relevant background. My capstone provides stronger evidence of independent development than of collaboration, so I still need to strengthen the connection between that experience and collaborative decisions. In future projects, I want to make responsibilities, integration requirements, and feedback clearer so contributors can understand how their work fits together.",
"Outcome 2 concerns professional oral, written, and visual communication. My code review and enhancement narratives explain the original application, planned improvements, and completed changes. Writing for an instructor or developer requires details about design and verification, while communicating with a user requires a clear explanation of what the application is doing. For example, the indexing interface explains why scanning is waiting and what the user can do next. I learned that accurate technical work becomes more useful when people can understand its results and limitations."
]),
('Algorithms Software Engineering and Databases', [
"The capstone brought software engineering, algorithms, and database concepts together in one application. In the software-design enhancement, I separated search coordination, navigation loading, and sidebar state into focused components. I also replaced Boolean file-operation results with explicit success, cancellation, and failure outcomes. This made the components easier to test and allowed the interface to distinguish a canceled operation from an error.",
"For Outcome 3, evaluating computing solutions through algorithmic principles and design tradeoffs, I improved file discovery and folder-size calculation. Explicit stacks and queues replaced arbitrary traversal-depth limits. A bottom-up pass calculates folder totals in O(n) time by adding each subtree total to its parent once. Cached layouts reduce repeated work in the disk-usage viewer. These choices also introduce costs: deeper traversal can require more work, and caches need memory limits and rules that prevent stale results. I learned to evaluate completeness and responsiveness together.",
"Outcome 4 concerns using established techniques and tools to deliver useful computing solutions. I applied C#, WPF, automated regression tests, and SQLite to Clearspace's actual workflows. The database enhancement replaced JSON tag storage with related tables, constraints, indexes, and transactions. The migration retains the original JSON, and bulk changes either commit together or roll back. This work helped me understand how database design supports behavior users depend on, rather than functioning as an isolated technical exercise."
]),
('Security and Professional Responsibility', [
"Outcome 5 concerns anticipating vulnerabilities and protecting data and resources. I applied that mindset through parameterized database queries, input validation, safe handling of folder links, and tests for failure cases. The file-protection enhancement uses authenticated encryption and password-derived protection without storing plaintext passwords or unwrapped file keys in the database. Tests cover wrong passwords, tampering, and failures around file replacement.",
"I learned to consider recovery as part of security. A database transaction cannot also commit a Windows file replacement, so the database and filesystem can be left in different states after an interruption. Documenting that boundary helped me avoid overstating the protection offered by the feature. The milestone implementation also has explicit limits, including supported local NTFS files and a 64 MiB per-file limit. Passing tests provides evidence for the cases exercised; it does not replace an independent security review."
]),
('Introduction to the Portfolio', [
"The portfolio presents Clearspace as one artifact enhanced across three categories. The software-design work establishes clearer responsibilities and more reliable operation feedback. The algorithms work improves discovery, aggregation, and the display of large file collections. The database work preserves metadata and supports file protection. Together, they demonstrate how maintainability, efficiency, persistence, and security affect the same user experience.",
"The original and enhanced source, code review, and narratives let readers compare my starting point with the completed enhancements. They also explain challenges and the limits of verification. Preparing this portfolio has given me a clearer picture of my strengths in application development and of the areas I need to develop further, especially collaboration and evaluating software in varied real-world environments. I want to carry that approach into software and game development as I continue building experience."
])]

doc = Document()
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(0.85)
sec.left_margin=sec.right_margin=Inches(1)
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
for name in ('Normal','Title','Subtitle','Heading 1'):
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.0
normal.paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(20)
doc.styles['Heading 1'].font.size=Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.add_paragraph('CS 499 Professional Self Assessment', 'Title')
doc.add_paragraph('Payton Castle\nOctober 11, 2026', 'Subtitle')
for heading, paragraphs in sections:
    heading_p=doc.add_heading(heading,1)
    if heading=='Algorithms Software Engineering and Databases':
        heading_p.paragraph_format.page_break_before=True
    for text in paragraphs:
        p=doc.add_paragraph(text)
        if text.startswith('I entered'):
            doc.add_comment(p.runs, 'Before final submission, add a specific coursework example outside Clearspace. Identify the course, what you did, and what it taught you. A course project has not yet been supplied.',author='Draft review')
        if text.startswith('For Outcome 1'):
            doc.add_comment(p.runs, 'Strengthen this outcome with one actual Unity collaboration decision: how you coordinated the music contribution, considered feedback, or resolved an integration issue. Add instructor feedback and resulting revisions only if they occurred.',author='Draft review')

# Include a descriptive, clickable source link for readers of the Word document.
p=doc.add_paragraph()
r=doc.part.relate_to('https://github.com/Friday7352/Clearspace','http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),r)
run=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'0563C1'); props.append(color); run.append(props)
t=OxmlElement('w:t'); t.text='View the Clearspace source repository'; run.append(t); h.append(run); p._p.append(h)
doc.core_properties.author='Payton Castle'
doc.core_properties.title='CS 499 Professional Self Assessment'
doc.core_properties.subject='Professional growth and introduction to the CS 499 ePortfolio'
doc.save(out)
print(out.resolve())
print('Body word count',sum(len(p.split()) for _,ps in sections for p in ps))
