from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

source = Path('C:/Users/Payton/OneDrive/Desktop/CS499/Module 3/Module 3 Journal.docx')
original = Document(source)
doc = Document()
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.8)
for key in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[key]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(8)
doc.styles['Title'].font.size = Pt(21)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
for key in ['Heading 1', 'Heading 2']:
    doc.styles[key].font.size = Pt(13)
    doc.styles[key].paragraph_format.space_before = Pt(10)
    doc.styles[key].paragraph_format.space_after = Pt(5)
doc.add_paragraph('CS 499 Career Reflection and Status Checkpoints', 'Title')
doc.add_paragraph('Payton Castle\nCS 499 Computer Science Capstone\nSeptember 26, 2026')
sections = [
('Career Plans',
 'When I started the Computer Science program, I thought it would be cool to make software and games. That is still the general direction I want to pursue. I have not settled on one specific role, but I am still interested in creating applications and games that people can use. Working on Clearspace has helped confirm that interest and given me a practical project to keep improving.'),
('How My Thinking Has Evolved',
 'My original interest was mostly in the idea of making something. Through my coursework and Clearspace, I have started to see more of what it takes to make software useful and dependable. For example, adding a disk usage viewer also meant dealing with slow performance, missing results in deep folders, and confusing indexing feedback. These experiences have made me think more about testing, usability, and how the application behaves with large amounts of data. My overall career direction has stayed similar, but I now have a clearer picture of the work involved.'),
('Career Research and Further Education',
 'I have not really explored specific jobs, certifications, or graduate programs yet, so career research has not had much influence on my plans so far. I am still exploring what direction would fit me best. A useful next step would be to compare software development and game development positions, look at the skills they require, and see how my projects relate to those expectations. I have not made a decision about an advanced degree or certification after my undergraduate degree.'),
('Course Outcomes and Remaining Work',
 'So far, my strongest evidence is in sound computing practices, professional communication, and algorithmic problem solving. For Artifact 1, I separated search, navigation, sidebar, and file-operation responsibilities and improved success, cancellation, and failure feedback. For Artifact 2, I removed arbitrary traversal-depth limits, added efficient bottom-up folder-size calculations, improved per-drive index recovery, and reused cached layouts. The latest full Release test run for that work passed 158 tests. The narrative and indexing explanations also help communicate the reasons for these decisions and their limitations.'),
]
for heading, text in sections:
    doc.add_paragraph(heading, 'Heading 1')
    doc.add_paragraph(text)
doc.add_paragraph('Validation, safe handling of folder links, and clearer error reporting support progress toward the security outcome. I still need to complete the database enhancement, gather stronger evidence for collaboration, and refine the final ePortfolio using instructor feedback. The database plan is to move tag storage from JSON to SQLite and test migration, constraints, and transactions. Real-drive checks and a review of the outcome mapping remain important before I consider the whole portfolio finished.')
doc.add_page_break()
doc.add_paragraph('Status Checkpoints for All Categories', 'Heading 1')
rows = [[c.text for c in row.cells] for row in original.tables[0].rows]
updates = {
1: 'Clearspace file index, search workflow, and disk usage analyzer.',
2: 'Completed: deep traversal, bottom-up folder totals, per-drive index recovery, cached layouts, and clearer indexing feedback.',
3: 'Implementation and Milestone Three Word narrative are complete. Refresh the code package for instructor review.',
4: 'Core enhancement complete. Latest full suite: 158 tests passed. Real-drive checks and instructor feedback remain.',
5: 'No. Publish after instructor review and any needed revisions.',
6: 'In progress. Algorithm tradeoffs and test evidence are documented; final outcome mapping and portfolio presentation remain.'
}
for i, text in updates.items(): rows[i][2] = text
table = doc.add_table(rows=0, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.18, 1.86, 1.92, 1.94]
for col, width in zip(table.columns, widths): col.width = Inches(width)
for i, values in enumerate(rows):
    cells = table.add_row().cells
    for j, (cell, text) in enumerate(zip(cells, values)):
        cell.width = Inches(widths[j])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = text
        props = cell._tc.get_or_add_tcPr()
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'E6E6E6' if i == 0 else ('F5F5F5' if i % 2 == 0 else 'FFFFFF'))
        props.append(shade)
        borders = OxmlElement('w:tcBorders')
        for edge in ['top', 'left', 'bottom', 'right']:
            element = OxmlElement('w:' + edge)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: element.set(qn('w:' + key), value)
            borders.append(element)
        props.append(borders)
        margins = OxmlElement('w:tcMar')
        for edge in ['top', 'left', 'bottom', 'right']:
            element = OxmlElement('w:' + edge)
            element.set(qn('w:w'), '100')
            element.set(qn('w:type'), 'dxa')
            margins.append(element)
        props.append(margins)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            for run in p.runs:
                run.font.size = Pt(10)
                run.bold = i == 0 or j == 0
    trpr = table.rows[-1]._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:cantSplit'))
    if i == 0: trpr.append(OxmlElement('w:tblHeader'))
doc.add_paragraph('Feedback Requested', 'Heading 2')
doc.add_paragraph('I would appreciate feedback on whether Artifact 2 clearly demonstrates algorithmic tradeoffs and whether the tests and narrative provide enough evidence for the algorithms and data structures outcome. I would also welcome guidance on the planned database migration before implementing it.')
doc.core_properties.author = 'Payton Castle'
doc.core_properties.title = 'CS 499 Career Reflection and Status Checkpoints'
output = Path(__file__).parent / 'Payton_Castle_Module_3_Journal_Updated.docx'
doc.save(output)
print(output)
