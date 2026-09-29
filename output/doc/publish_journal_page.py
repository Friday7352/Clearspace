from pathlib import Path
from html import escape
from shutil import copyfile
from docx import Document

site = Path(__file__).resolve().parents[1] / 'eportfolio-publish'
source = Path(__file__).parent / 'Payton_Castle_Module_3_Journal_Updated.docx'
target = site / 'assets/journals' / source.name
target.parent.mkdir(parents=True, exist_ok=True)
copyfile(source, target)
d = Document(source)
index = (site / 'index.html').read_text(encoding='utf-8')
css = index.split('<style>', 1)[1].split('</style>', 1)[0]
parts = []
for element in d.element.body:
    if element.tag.endswith('}p'):
        p = next((p for p in d.paragraphs if p._p is element), None)
        if p is None or not p.text or p.style.name == 'Title': continue
        if p.text.startswith('Payton Castle\n'): continue
        tag = 'h2' if p.style.name.startswith('Heading') else 'p'
        parts.append(f'<{tag}>{escape(p.text)}</{tag}>')
    elif element.tag.endswith('}tbl'):
        table = next(t for t in d.tables if t._tbl is element)
        parts.append('<div class="table-scroll" tabindex="0" role="region" aria-label="Artifact status checkpoints"><table><caption>Enhancement status as of September 26, 2026</caption><thead>')
        for i, row in enumerate(table.rows):
            if i == 1: parts.append('</thead><tbody>')
            parts.append('<tr>')
            for j, cell in enumerate(row.cells):
                tag = 'th' if i == 0 or j == 0 else 'td'
                scope = ' scope="col"' if i == 0 else ' scope="row"' if j == 0 else ''
                parts.append(f'<{tag}{scope}>{escape(cell.text)}</{tag}>')
            parts.append('</tr>')
        parts.append('</tbody></table></div>')
page = '''<!DOCTYPE html>
<html lang="en"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Career Reflection and Status Checkpoints | Payton</title>
<style>''' + css + '''
main a, header a { color: var(--accent); }
.table-scroll { overflow-x: auto; margin: 20px 0 32px; }
table { width: 100%; min-width: 680px; border-collapse: collapse; font-size: .9rem; }
caption { text-align: left; color: var(--muted); margin-bottom: 12px; }
th, td { border: 1px solid var(--border); padding: 12px; text-align: left; vertical-align: top; }
th { background: var(--panel); }
tbody tr:nth-child(even) td { background: var(--panel); }
main h2 { margin-top: 32px; }
@media(max-width: 600px) { header { padding-top: 32px; } header h1 { font-size: 1.9rem; } }
</style></head><body><header>
<a href="index.html#journals">&larr; Back to ePortfolio</a>
<h1>Career Reflection and Status Checkpoints</h1>
<p class="tagline">Payton Castle &middot; CS 499 &middot; September 26, 2026</p>
<p><a href="assets/journals/''' + source.name + '''">Download this journal as a Word document</a></p>
</header><main>''' + '\n'.join(parts) + '''</main>
<footer>&copy; 2026 Payton &middot; CS 499 Capstone ePortfolio</footer></body></html>'''
(site / 'career-reflection.html').write_text(page, encoding='utf-8')
index = index.replace('  <a href="#code-review">Code Review</a>', '  <a href="#code-review">Code Review</a>\n  <a href="#journals">Journals</a>')
section = '''  <section id="journals">
    <h2>Journals and Progress</h2>
    <div class="card">
      <h3>Career Reflection and Status Checkpoints</h3>
      <p>September 26, 2026 &mdash; My interest in software and game development,
      what I have learned through Clearspace, and progress across the three enhancement categories.</p>
      <ul class="artifact-links">
        <li><a class="repo-link" href="career-reflection.html">Read the journal &rarr;</a></li>
        <li><a class="repo-link" href="assets/journals/''' + source.name + '''">Download the journal (DOCX)</a></li>
      </ul>
    </div>
  </section>

'''
index = index.replace('</main>', section + '</main>')
(site / 'index.html').write_text(index, encoding='utf-8')
print('Prepared journal page, Word download, and homepage links.')
