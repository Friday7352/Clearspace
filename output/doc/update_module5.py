from pathlib import Path
from docx import Document
import re, shutil, html, urllib.request, json

root=Path(r'E:/GithubRepos/File Explorer Remake')
src=Path(r'C:/Users/Payton/OneDrive/Desktop/CS499/Module 5/Module 5 Journal.docx')
out=root/'output/doc/Payton_Castle_Module_5_Journal_Updated.docx'
d=Document(src)
texts={
0:'CS 499 Disruptive Technologies and Status Checkpoints',
2:'Technologies I Chose',
3:'I chose generative AI for software development and extended reality (XR), which includes virtual reality and augmented reality. Both are game changers because they change how software is created or how people interact with it. They also connect to the software and games I want to build.',
4:'Identification and Description of Each Technology',
5:'Generative AI creates content, including code, from prompts and context. AI coding tools can suggest implementations, explain existing code, and assist with reviews. GitHub Copilot is one example (GitHub, n.d.). This breaks the older pattern in which developers manually wrote most implementation details and creates a trend toward AI-assisted development. Developers still need to decide whether the output is correct and suitable.',
6:'XR combines immersive virtual environments with digital content placed in the physical world. VR replaces the view with a simulated environment, while AR adds digital objects to the real one. XR breaks the screen-and-mouse interaction pattern by using spatial views and tracked movement. It has encouraged immersive games, simulations, and shared environments. OpenXR provides a common API across XR devices, reducing the need for separate device-specific integrations (Khronos Group, n.d.).',
7:'Likely Impacts on Computer Science',
8:'AI-assisted development changes the industry workflow by making code generation and automated review part of everyday development. I expect more attention to requirements, architecture, testing, and security as developers evaluate generated changes. Faster first drafts can help, but incorrect output creates extra review work. The tests I wrote for Clearspace show why checking behavior matters regardless of how code is produced.',
9:'XR increases the need for real-time graphics, spatial input, accessible interaction, and efficient networking. Developers must balance image quality with frame time and user comfort. Shared APIs can make applications easier to move between devices. My work on Clearspace performance has given me practice with responsiveness and resource tradeoffs, although immersive applications would require additional graphics and interaction skills.',
10:'Impacts on Humans Communities and the World',
11:'AI coding tools could help individuals and small teams build useful software sooner and make programming easier to approach. Communities could benefit from tools tailored to local needs. However, mistakes, biased output, and insecure code can affect many users when software is distributed. Workers may spend more time reviewing automated work, and people entering the field will still need opportunities to build their own problem-solving skills.',
12:'XR could support remote social interaction, immersive education, and training where practicing in the real world is expensive or dangerous. These uses may connect people across distance and give communities new ways to share experiences. Hardware cost, motion sickness, and accessibility barriers can exclude users. Tracking data and harassment in shared environments also require careful privacy and moderation choices.',
13:'Likely Impacts on My Career',
14:'Both technologies fit my goal of building software and games. I expect AI tools to be part of my development workflow, so I need to understand their suggestions, test changes, and protect user data. Clearspace gave me practice with validation, regression testing, and explaining technical limits. I want to use AI assistance while continuing to develop those skills.',
15:'XR interests me because I build VRChat worlds as a hobby, and the C# I use in Clearspace carries over to Unity. Working with XR would push me to learn more about 3D interaction and performance. Comparing software and game development roles will help me decide which skills to develop next.',
17:'I have demonstrated Outcome 2 through the enhancement narratives and technical explanations for different audiences. Outcome 3 is supported by my algorithm choices and tradeoffs: deep traversal, bottom-up folder totals, cached layouts, and per-drive index recovery. The Artifact 2 verification passed 158 Release tests. Outcome 4 is supported by using C#, WPF, SQLite, and automated tests to deliver working enhancements across all three categories.',
18:'I have also demonstrated Outcome 5 through parameterized database queries, input validation, transactions, and password-based file protection. Artifact 3 moved tag storage from JSON to SQLite while preserving the original JSON during migration. Its October 1 verification passed 251 Release tests. These are results from the milestone snapshots, not a claim about every later change to Clearspace.',
19:'Outcome 1 remains the area where I need stronger evidence. I built Clearspace mainly on my own, so I still need to document how instructor or peer feedback improves the work and supports collaborative decisions. My portfolio now connects the artifacts to the outcomes, but final instructor review and the code-review video remain. Additional real-drive checks would strengthen performance evidence; they are not implied by the automated test counts.',
23:'Please tell me whether the artifacts and narratives provide sufficient evidence for Outcomes 2 through 5. I would also appreciate guidance on documenting collaboration for Outcome 1. The code-review video will be added when it is uploaded.',
}
for i,s in texts.items():
 p=d.paragraphs[i]
 if p.runs:
  p.runs[0].text=s
  for r in p.runs[1:]: r.text=''
 else: p.add_run(s)
t=d.tables[0]
rows={
3:['Packages and narratives prepared; instructor submission status needs confirmation.']*3,
4:['Core enhancement complete; the portfolio documents architecture and 87 passing milestone tests. Instructor revisions remain if requested.','Core enhancement complete; 158 passing milestone tests. Real-drive checks would strengthen evidence.','Core enhancement complete; 251 passing milestone tests. Filesystem and database recovery limits are documented.'],
5:['Narrative and source links included in the ePortfolio.']*3,
6:['Category content and outcome mapping complete. Final portfolio awaits the code-review video and instructor review.']*3,
}
for row,values in rows.items():
 for col,value in enumerate(values,1):
  p=t.cell(row,col).paragraphs[0]
  if p.runs:
   p.runs[0].text=value
   for r in p.runs[1:]: r.text=''
  else: p.add_run(value)
d.add_paragraph('References',style=d.paragraphs[22].style)
d.add_paragraph('GitHub. (n.d.). GitHub Copilot. https://github.com/features/copilot')
d.add_paragraph('Khronos Group. (n.d.). OpenXR. https://www.khronos.org/openxr/')
d.save(out)

self_assessment='''<p>I entered the Computer Science program interested in making software and games. Building Clearspace, a Windows file manager in C# and WPF, helped me connect that interest to the practical work of designing, testing, and maintaining an application. I learned to evaluate success through user feedback, reliable behavior, and clear explanations of limitations.</p>
<p>In software design and engineering, I separated search, navigation, sidebar, and file-operation responsibilities and improved feedback for success, cancellation, and failure. In algorithms and data structures, I replaced depth-limited traversal, calculated folder totals in a bottom-up pass, and reused cached layouts. In databases, I replaced JSON tag storage with SQLite constraints, indexes, and transactions, then added records that support password-protected files. Each enhancement addresses a different part of the same application.</p>
<p>The milestone verification grew from 87 Release tests for the software-design enhancement to 158 for algorithms and 251 for databases. These results support the submission snapshots; they do not guarantee performance on every drive or establish an independent security audit. The narratives explain the tradeoffs and the remaining boundaries, including separate filesystem and database commits.</p>
<p>My strongest demonstrated outcomes are professional communication, algorithmic problem solving, use of established computing tools, and a security mindset. Collaboration needs additional evidence from instructor or peer feedback, since I developed Clearspace mainly independently. I want to bring the habits of testing, documenting decisions, and protecting user data into a career in software or game development. Building VRChat worlds also motivates me to explore Unity and real-time 3D interaction.</p>'''
mapping='''<section id="outcomes"><h2>Course Outcome Evidence</h2>
<ul><li><strong>Outcome 1 — Collaborative environments:</strong> Clearspace was built mainly independently. Instructor or peer feedback and the resulting decisions still need to be documented to strengthen this outcome.</li>
<li><strong>Outcome 2 — Professional communication:</strong> The three enhancement narratives explain purpose, implementation choices, verification, and limitations. The portfolio presents these explanations alongside source and evidence.</li>
<li><strong>Outcome 3 — Algorithmic solutions and tradeoffs:</strong> The algorithms enhancement demonstrates traversal, bottom-up aggregation, layout caching, and memory versus responsiveness tradeoffs. The database narrative also discusses migration and recovery choices.</li>
<li><strong>Outcome 4 — Computing techniques and tools:</strong> C#, WPF, indexed data structures, SQLite, and regression tests turn the enhancement plans into working features.</li>
<li><strong>Outcome 5 — Security mindset:</strong> Parameterized queries, validation, relational constraints, authenticated encryption, and failure-case tests address threats to metadata and files. Documented limits distinguish verified behavior from further security work.</li></ul></section>'''
for folder in ['eportfolio-publish','portfolio-repository']:
 base=root/'output'/folder
 index=base/'index.html'
 s=index.read_text(encoding='utf-8')
 s=re.sub(r'<p><em>Coming soon.*?</em></p>',self_assessment,s,flags=re.S)
 s=s.replace('genuinely useful','useful')
 s=s.replace('<section id="artifact">',mapping+'\n<section id="artifact">')
 s=s.replace('<a href="#artifact">Artifact</a>','<a href="#outcomes">Course Outcomes</a>\n  <a href="#artifact">Artifact</a>')
 s=s.replace('<p><em>The video link will be added here.</em></p>','<p>The code-review video is awaiting upload. I will add the recording here when it is available.</p>')
 s=s.replace('<section id="journals">','<section id="journals">\n<h3>Disruptive Technologies and Status Checkpoints</h3><p>October 4, 2026 — Generative AI, extended reality, course outcomes, and progress across all three enhancements.</p><p><a class="repo-link" href="module-five-journal.html">Read the journal</a> · <a class="repo-link" href="assets/journals/Payton_Castle_Module_5_Journal_Updated.docx">Download the journal (DOCX)</a></p>')
 index.write_text(s,encoding='utf-8')
 shutil.copy2(out,base/'assets/journals'/out.name)
 css=re.search(r'<style>(.*?)</style>',(base/'career-reflection.html').read_text(encoding='utf-8'),re.S).group(1)
 body=[]
 for i,p in enumerate(d.paragraphs):
  if i in [0,1] or not p.text.strip(): continue
  if i==20:
   body.append('<h2>'+html.escape(p.text)+'</h2><div class="table-scroll"><table>')
   for ri,r in enumerate(t.rows):
    body.append('<tr>'+''.join(('<th>' if ri==0 or ci==0 else '<td>')+html.escape(c.text)+('</th>' if ri==0 or ci==0 else '</td>') for ci,c in enumerate(r.cells))+'</tr>')
   body.append('</table></div>')
  else:
   tag='h2' if i in [2,4,7,10,13,16,22,25] else 'p'
   body.append(f'<{tag}>{html.escape(p.text)}</{tag}>')
 (base/'module-five-journal.html').write_text('<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Disruptive Technologies | Payton Castle</title><style>'+css+'</style></head><body><header><a href="index.html#journals">Back to ePortfolio</a><h1>Disruptive Technologies and Status Checkpoints</h1><p>Payton Castle · CS 499 · October 4, 2026</p></header><main>'+''.join(body)+'</main></body></html>',encoding='utf-8')

try:
 response=urllib.request.urlopen('https://friday7352.github.io/',timeout=30)
 live=response.read().decode()
 (root/'output/doc/live-portfolio.html').write_text(live,encoding='utf-8')
 print('Live site:',response.status,'self-assessment placeholder:', 'Coming soon' in live,'video placeholder:', 'video link will' in live)
except Exception as e: print('Live site check:',e)
print('Saved',out)
